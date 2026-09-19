from __future__ import annotations

from io import BytesIO
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .models import Acta, ActaVoteEntry, ManualReviewQueue


HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)


def _style_header(ws, ncols, first_row=1):
    for col in range(1, ncols + 1):
        c = ws.cell(row=first_row, column=col)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = CENTER
        c.border = BORDER


def _autosize(ws, widths: dict | None = None):
    widths = widths or {}
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        max_len = widths.get(letter, 12)
        for cell in col:
            try:
                v = cell.value
                if v is not None:
                    max_len = max(max_len, min(60, len(str(v))))
            except Exception:
                pass
        ws.column_dimensions[letter].width = max_len + 2


def export_actas_xlsx(actas_qs) -> Workbook:
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Votos Detalle"
    headers = [
        "Acta ID", "Unique ID", "Tipo Acta", "Mesa", "Ubigeo",
        "Departamento", "Provincia", "Distrito",
        "Proceso Electoral", "Estado Acta",
        "Categoria", "Organizacion (cod)", "Organizacion (nombre)",
        "Candidato", "Cargo", "Ambito", "Orden",
        "Votos", "Origen Dato", "Confianza Campo",
        "Confianza Global Transcripcion",
    ]
    ws1.append(headers)
    _style_header(ws1, len(headers))

    actas_ids = list(actas_qs.values_list("id", flat=True))
    entries = ActaVoteEntry.objects.filter(acta_id__in=actas_ids).select_related(
        "acta", "acta__election_process", "acta__department",
        "acta__province", "acta__district", "organization",
        "transcription",
    )

    row_num = 2
    for e in entries:
        a = e.acta
        t = e.transcription
        ws1.append([
            a.id, a.unique_id, a.acta_type, a.table_number,
            a.district.ubigeo if a.district else "",
            a.department.name if a.department else "",
            a.province.name if a.province else "",
            a.district.name if a.district else "",
            a.election_process.name if a.election_process else "",
            a.status,
            e.entry_category,
            e.organization.code if e.organization else "",
            e.organization.short_name if e.organization else e.candidate_name,
            e.candidate_name, e.position, e.scope, e.sort_order,
            e.votes, e.data_origin, e.field_confidence,
            t.global_confidence if t else 0.0,
        ])
        row_num += 1
    _autosize(ws1)

    ws2 = wb.create_sheet("Resumen por Acta")
    from django.db.models import Sum, Count
    headers2 = [
        "Acta ID", "Unique ID", "Tipo", "Proceso", "Departamento", "Provincia", "Distrito", "Mesa",
        "Estado", "Votos Candidatos", "Blanco", "Nulo", "Impugnado",
        "Total Calculado", "Total Consignado", "Diferencia", "Transcripciones",
        "Confianza Global Max",
    ]
    ws2.append(headers2)
    _style_header(ws2, len(headers2))
    for a in actas_qs.select_related("election_process", "department", "province", "district"):
        vs = a.vote_entries
        votables = vs.filter(entry_category__in={
            "organizacion", "gobernador", "consejero", "provincial", "distrital",
        }).aggregate(s=Sum("votes"))["s"] or 0
        blanco = vs.filter(entry_category="blanco").aggregate(s=Sum("votes"))["s"] or 0
        nulo = vs.filter(entry_category="nulo").aggregate(s=Sum("votes"))["s"] or 0
        impug = vs.filter(entry_category="impugnado").aggregate(s=Sum("votes"))["s"] or 0
        total = vs.filter(entry_category="total").aggregate(s=Sum("votes"))["s"] or 0
        calc = votables + blanco + nulo + impug
        diff = (total - calc) if total else None
        conf_max = 0.0
        for t in a.transcriptions.all():
            if t.global_confidence > conf_max:
                conf_max = t.global_confidence
        ws2.append([
            a.id, a.unique_id, a.acta_type,
            a.election_process.name if a.election_process else "",
            a.department.name if a.department else "",
            a.province.name if a.province else "",
            a.district.name if a.district else "",
            a.table_number, a.status,
            votables, blanco, nulo, impug, calc, total, diff,
            a.transcriptions.count(), conf_max,
        ])
    _autosize(ws2)

    ws3 = wb.create_sheet("Actas Observadas")
    headers3 = ["Rev ID", "Acta ID", "Estado", "Asignado", "Razon", "Codigos",
                "Abierto", "Cerrado", "Resolucion"]
    ws3.append(headers3)
    _style_header(ws3, len(headers3))
    revs = ManualReviewQueue.objects.filter(acta_id__in=actas_ids).select_related("acta", "assigned_to", "closed_by")
    for r in revs:
        ws3.append([
            r.id, r.acta_id, r.status,
            r.assigned_to.username if r.assigned_to else "",
            r.reason,
            ", ".join(r.reason_codes) if isinstance(r.reason_codes, list) else r.reason_codes,
            str(r.opened_at),
            str(r.closed_at) if r.closed_at else "",
            r.resolution_notes,
        ])
    _autosize(ws3)

    ws4 = wb.create_sheet("Metricas")
    ws4.append(["Métrica", "Valor"])
    _style_header(ws4, 2)
    total_actas = actas_qs.count()
    ws4.append(["Total actas", total_actas])
    for status, label in Acta._meta.get_field("status").choices:
        ws4.append([f"Actas - {label}", actas_qs.filter(status=status).count()])
    ws4.append(["Total entradas de votos", ActaVoteEntry.objects.filter(acta_id__in=actas_ids).count()])
    ws4.append(["Actas en revisión", revs.count()])
    _autosize(ws4, {"A": 30, "B": 18})

    return wb


def export_review_xlsx() -> Workbook:
    qs = Acta.objects.filter(status="observado").order_by("id")
    return export_actas_xlsx(qs)
